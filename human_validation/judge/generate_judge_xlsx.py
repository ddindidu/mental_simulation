#!/usr/bin/env python3
"""
Sample 4 plain-style transcripts per disorder (2 from gemini-3.8-flash /
claude-sonnet-5, 2 from llama-3.3-70b-instruct / qwen3-235b) and render each
as a judge-facing .xlsx in human_validation/judge/.
"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent.parent
RUN = BASE / "saved" / "run_batch_20260912"
PROFILES_DIR = BASE / "data" / "v2_final_profiles"
KG_DISORDER = BASE / "mentalbench" / "resources" / "knowledge_graph" / "EN" / "disorder.json"
OUT_DIR = Path(__file__).resolve().parent

PATIENT_ROOT = "gpt-5.6-terra"  # patient-simulator model whose logs we sample from
SET_A = ["gemini-3.8-flash", "claude-sonnet-5"]
SET_B = ["llama-3.3-70b-instruct", "qwen3-235b"]
N_PER_DISORDER = 4
SEED = 42

FONT_NAME = "Calibri"


def load_json(p: Path):
    with open(p, encoding="utf-8") as f:
        return json.load(f)


DISORDER_NAMES = {k: v["name"] for k, v in load_json(KG_DISORDER).items()}


def dname(did: str) -> str:
    return DISORDER_NAMES.get(did, did) if did else did


def results_dir(doctor: str) -> Path:
    return RUN / "results" / PATIENT_ROOT / PATIENT_ROOT / doctor


def analysis_dir(doctor: str) -> Path:
    return RUN / "analysis" / PATIENT_ROOT / PATIENT_ROOT / doctor


# ── Data loading for one case ───────────────────────────────────────────────

class CaseData:
    def __init__(self, profile_id: str, doctor: str):
        self.profile_id = profile_id
        self.doctor = doctor
        self.disorder_code = profile_id.split("_")[0]
        self.log_file = f"{profile_id}_plain"

        self.profile = load_json(PROFILES_DIR / self.disorder_code / f"{profile_id}.json")
        self.transcript = load_json(results_dir(doctor) / f"{profile_id}_plain_result.json")

        eff_list = load_json(results_dir(doctor) / "efficiency_eval.json")
        self.efficiency = next(e for e in eff_list if e["log_file"] == self.log_file)

        dr_list = load_json(results_dir(doctor) / "diagnostic_reasoning_eval.json")
        self.diagnostic_reasoning = next(e for e in dr_list if e["log_file"] == self.log_file)

        turn_eval_list = load_json(analysis_dir(doctor) / "turn_eval.json")
        self.turn_evals = [e for e in turn_eval_list if e["log_file"] == self.log_file]

        qe_list = load_json(analysis_dir(doctor) / "question_eval_semantic.json")
        qe_entry = next((e for e in qe_list if e["log_file"] == self.log_file), None)
        self.question_turns = qe_entry["turns"] if qe_entry else []

    @staticmethod
    def _mean(vals):
        vals = [v for v in vals if v is not None]
        return round(sum(vals) / len(vals), 4) if vals else None

    def scores_table(self):
        return [
            ("Jaccard (mean over turns)", self._mean(t.get("jaccard") for t in self.turn_evals)),
            ("Precision (mean over turns)", self._mean(t.get("precision") for t in self.turn_evals)),
            ("Recall (mean over turns)", self._mean(t.get("recall") for t in self.turn_evals)),
            ("IAS (mean over turns)", self._mean(t.get("ias") for t in self.question_turns)),
            ("ECR (mean over turns)", self._mean(t.get("ecr") for t in self.question_turns)),
            ("Turn Count", self.efficiency.get("turn_count")),
            ("Turns to First Confident Narrowing", self.efficiency.get("time_to_first_confident_narrowing")),
            ("Final Accuracy", self.efficiency.get("final_accuracy")),
            ("Diagnostic Reasoning Ability (overall_score)", self.diagnostic_reasoning.get("overall_score")),
        ]

    def patient_profile_text(self):
        cp = self.profile["clinical_profiles"]
        pi = self.profile["patient_info"]
        feats = cp.get("sampled_features", cp)
        crit = feats.get("sampled_diagnostic_criteria", {})
        lines = [
            f"Disorder: {cp.get('disease_name')} ({cp.get('disease_code')})",
            f"Demographics: {pi.get('demographics')}",
            f"Severity: {pi.get('severity')}",
            f"Stressor event: {pi.get('stressor_event')}",
            f"Duration: {crit.get('duration')}",
            f"Additional requirements: {', '.join(crit.get('additional_requirements', []) or [])}",
            f"Chief complaint symptom code: {pi.get('chief_complaint_symptom_code')}",
            "",
            f"Background: {pi.get('stressor')}",
        ]
        return "\n".join(str(l) for l in lines)

    @staticmethod
    def _candidate_str(candidate_set: dict) -> str:
        parts = []
        for tier, label in (
            ("high_likely", "High-likely"),
            ("moderate_likely", "Moderate-likely"),
            ("low_likely", "Low-likely"),
            ("excluded", "Excluded"),
        ):
            ids = candidate_set.get(tier) or []
            if not ids:
                continue
            names = "; ".join(f"{dname(i)} ({i})" for i in ids)
            parts.append(f"{label}: {names}")
        return " | ".join(parts) if parts else "(no differential recorded for this turn)"

    def turns(self):
        for t in self.transcript["turns"]:
            yield (
                t["turn"],
                t["doctor_question"],
                t["patient_response"],
                self._candidate_str(t.get("candidate_set", {})),
            )

    def final_diagnosis_text(self):
        eff = self.efficiency
        return (
            f"Doctor's Final Diagnosis: {dname(eff.get('final_diagnosis_id'))} "
            f"({eff.get('final_diagnosis_id')}, ICD: {eff.get('final_diagnosis')})"
        )

    def checklist_rows(self):
        dr = self.diagnostic_reasoning
        rows = [
            f"Duration requirement — score: {dr.get('duration_score')}",
            f"Functional impairment requirement — score: {dr.get('functional_impairment_score')}",
            f"Additional requirements — score: {dr.get('additional_requirements_score')}",
        ]
        for grp in dr.get("symptom_criterion_evaluations") or []:
            matched = "; ".join(grp.get("matched_symptom_descriptions") or []) or "(none)"
            missing = "; ".join(grp.get("missing_required_symptoms") or []) or "(none)"
            rows.append(
                f"Checklist group: {grp.get('group')} "
                f"(relation={grp.get('relation')}, required={grp.get('min_count_required')}, "
                f"matched_count={grp.get('valid_symptom_count')}, coverage={grp.get('symptom_coverage_score')})\n"
                f"  Matched evidence: {matched}\n"
                f"  Missing: {missing}"
            )
        return rows


# ── Sampling ─────────────────────────────────────────────────────────────

def list_profile_ids(disorder_code: str):
    return sorted(p.stem for p in (PROFILES_DIR / disorder_code).glob("*.json"))


def _has_case_files(profile_id: str, doctor: str) -> bool:
    disorder_code = profile_id.split("_")[0]
    log_file = f"{profile_id}_plain"
    return (
        (PROFILES_DIR / disorder_code / f"{profile_id}.json").exists()
        and (results_dir(doctor) / f"{profile_id}_plain_result.json").exists()
        and (results_dir(doctor) / "efficiency_eval.json").exists()
        and (results_dir(doctor) / "diagnostic_reasoning_eval.json").exists()
        and (analysis_dir(doctor) / "turn_eval.json").exists()
    )


def sample_cases():
    rng = random.Random(SEED)
    disorders = sorted(p.name for p in PROFILES_DIR.iterdir() if p.is_dir())
    samples = []
    for disorder in disorders:
        profiles = list_profile_ids(disorder)
        rng.shuffle(profiles)
        pool = iter(profiles)
        used_profiles = set()
        for doctor_set in (SET_A, SET_A, SET_B, SET_B):
            doctors = list(doctor_set)
            rng.shuffle(doctors)
            picked = None
            for profile_id in pool:
                if profile_id in used_profiles:
                    continue
                for doctor in doctors:
                    if _has_case_files(profile_id, doctor):
                        picked = (profile_id, doctor)
                        break
                if picked:
                    break
            if picked is None:
                print(f"  WARN: could not find a valid case for disorder {disorder} ({doctor_set})")
                continue
            used_profiles.add(picked[0])
            samples.append(picked)
    return samples


# ── XLSX rendering ───────────────────────────────────────────────────────

TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True)
SECTION_FONT = Font(name=FONT_NAME, size=13, bold=True)
LABEL_FONT = Font(name=FONT_NAME, size=11, bold=True)
BODY_FONT = Font(name=FONT_NAME, size=10)
SECTION_FILL = PatternFill("solid", fgColor="E8E8E8")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def write_section_header(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT
    for col in (1, 2):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    return row + 1


def build_workbook(case: CaseData) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transcript"

    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 90

    row = 1
    title_cell = ws.cell(
        row=row, column=1,
        value=f"Case ID: {case.profile_id}    |    Doctor Model: {case.doctor}",
    )
    title_cell.font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 2

    row = write_section_header(ws, row, "Automated Metrics")
    for label, value in case.scores_table():
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=row, column=2, value=value).font = BODY_FONT
        row += 1
    row += 1

    row = write_section_header(ws, row, "Patient Profile")
    cell = ws.cell(row=row, column=1, value=case.patient_profile_text())
    cell.font = BODY_FONT
    cell.alignment = WRAP_TOP
    row += 2

    row = write_section_header(ws, row, "Transcript")
    transcript_start_row = row
    for turn, question, response, candidates in case.turns():
        c = ws.cell(row=row, column=2, value=f"[Turn {turn} · Doctor Question] {question}")
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1

        c = ws.cell(row=row, column=1, value=f"[Turn {turn} · Patient Response] {response}")
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1

        c = ws.cell(row=row, column=2, value=f"[Turn {turn} · Doctor Differential] {candidates}")
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1
    row += 1

    row = write_section_header(ws, row, "Final Diagnosis & Diagnostic Checklist")
    c = ws.cell(row=row, column=1, value=case.final_diagnosis_text())
    c.font = LABEL_FONT
    c.alignment = WRAP_TOP
    row += 1

    for entry in case.checklist_rows():
        c = ws.cell(row=row, column=1, value=entry)
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1

    ws.freeze_panes = f"A{transcript_start_row}"
    return wb


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    samples = sample_cases()
    print(f"Sampling {len(samples)} cases ({N_PER_DISORDER} per disorder)")
    written = 0
    for profile_id, doctor in samples:
        try:
            case = CaseData(profile_id, doctor)
        except Exception as e:
            print(f"  SKIP {profile_id} / {doctor}: {e}")
            continue
        wb = build_workbook(case)
        out_path = OUT_DIR / f"judge_{profile_id}__{doctor}.xlsx"
        wb.save(out_path)
        written += 1
        print(f"  wrote {out_path.name}")
    print(f"Done: {written}/{len(samples)} workbooks written to {OUT_DIR}")


if __name__ == "__main__":
    main()
