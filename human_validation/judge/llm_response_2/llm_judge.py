#!/usr/bin/env python3
"""
LLM-as-judge scoring for the expert_response_2 validation round, using
OpenAI "gpt-6-luna" as judge model.

Scores the same 3 dimensions as
human_validation/judge/expert_response_2/전문가 평가 가이드라인_2.docx:
  1. Diagnostic Hypothesis Quality  (h) — per turn t >= 1.
  2. Diagnostic Question Quality    (q) — per turn t that has a question
     (t = 1 .. T-1). Turns follow utils/turn_policy.py: T0 is the opening
     question (not a turn, not evaluated); turn t = patient response t, the
     DDx after it, and the doctor's next question; the last turn has none.
  3. Diagnostic Evidence Sufficiency (evidence_sufficiency) — once per case,
     judged against the DOCTOR'S OWN final diagnosis (not ground truth) —
     i.e. the "_pred" semantics established in eval/evaluate_diagnostic_reasoning.py,
     since the guideline explicitly compares doctor's final diagnosis <-> DSM-5
     criteria, not patient profile <-> DSM-5.

One LLM call per case (not per turn) to keep token usage down: the whole
transcript + checklist is sent once, and the model returns per-turn +
episode-level scores in a single compact JSON reply. Candidate/differential
codes are sent as bare disease IDs (D0xx) with a legend, not repeated full
names, and checklist evidence is sent as short symptom names, not full DSM
descriptions — both to cut input tokens without losing what the rubric
actually needs.

Usage:
  python human_validation/judge/llm_response_2/llm_judge.py --print-prompt D001-1   # inspect one filled prompt, no API call
  python human_validation/judge/llm_response_2/llm_judge.py --case D001-1           # run just one case
  python human_validation/judge/llm_response_2/llm_judge.py                        # run all 46 cases
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
from pathlib import Path

THIS_DIR = Path(__file__).resolve().parent
JUDGE_DIR = THIS_DIR.parent
REPO_ROOT = JUDGE_DIR.parent.parent
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(JUDGE_DIR))

from dotenv import load_dotenv  # noqa: E402
load_dotenv(REPO_ROOT / ".env")

from generate_judge_xlsx import CaseData, DISORDER_NAMES, results_dir  # noqa: E402
from utils.turn_policy import policy_turns  # noqa: E402

# Recorded in every output file so readers know how "t" is numbered. Outputs
# written before this field existed numbered q by the patient response that
# FOLLOWS the question (q at t = the question of policy turn t - 1).
TURN_NUMBERING = "policy"

MODEL = "gpt-6-luna"
EXPERT_RESPONSE_2_DIR = JUDGE_DIR / "expert_response_2"
OUT_DIR = THIS_DIR


# ── Case discovery (same 46 cases as expert_response_2's xlsx sheets) ──────

def _normalize_doctor_label(label: str) -> str | None:
    """The xlsx header's doctor label has inconsistent spelling ('Clause-sonnet-5',
    'LLaMa-3,3-70b-instruct'), and every doctor has logs for every profile_id
    (full batch x all profiles), so it must be parsed from the label — probing
    'which doctor dir has this profile' doesn't disambiguate, it always matches."""
    m = label.lower()
    if "qwen" in m:
        return "qwen3-235b"
    if "claude" in m or "clause" in m:
        return "claude-sonnet-5"
    if "gemini" in m:
        return "gemini-3.8-flash"
    if "llama" in m:
        return "llama-3.3-70b-instruct"
    return None


def load_case_list() -> list[tuple[str, str, str]]:
    """[(sheet_name, profile_id, doctor), ...]"""
    import openpyxl

    path = sorted(glob.glob(str(EXPERT_RESPONSE_2_DIR / "doctor_*1.xlsx")))[0]
    wb = openpyxl.load_workbook(path, data_only=True)
    cases = []
    for sheet in wb.sheetnames:
        header = wb[sheet].cell(row=1, column=2).value
        if not header or "/" not in str(header):
            continue
        profile_id, doctor_label = str(header).split("/", 1)
        profile_id = profile_id.strip()
        doctor = _normalize_doctor_label(doctor_label.strip())
        if doctor is None or not (results_dir(doctor) / f"{profile_id}_plain_result.json").exists():
            print(f"  [warn] {sheet}: couldn't resolve doctor from label {doctor_label!r}, skipping", file=sys.stderr)
            continue
        cases.append((sheet, profile_id, doctor))
    return cases


# ── Compact prompt construction ─────────────────────────────────────────────

def _short_name(desc: str) -> str:
    """'Detail Inattention Errors: Often fails to ...' -> 'Detail Inattention Errors'"""
    return desc.split(":", 1)[0].strip()


def _tier_codes(candidate_set: dict) -> str:
    parts = []
    for tier, tag in (("high_likely", "H"), ("moderate_likely", "M"), ("low_likely", "L")):
        ids = candidate_set.get(tier) or []
        if ids:
            parts.append(f"{tag}:{','.join(ids)}")
    return " ".join(parts) if parts else "(none)"


def _legend(case: CaseData) -> str:
    codes = set()
    for t in case.transcript["turns"]:
        cs = t.get("candidate_set") or {}
        for tier in ("high_likely", "moderate_likely", "low_likely"):
            codes.update(cs.get(tier) or [])
    codes.add(case.efficiency.get("final_diagnosis_id") or "")
    codes.add(case.diagnostic_reasoning.get("ground_truth") or "")
    codes.discard("")
    return "; ".join(f"{c}={DISORDER_NAMES.get(c, c)}" for c in sorted(codes))


def _checklist_block(case: CaseData) -> str:
    dr = case.diagnostic_reasoning
    has_pred = dr.get("predicted_diagnosis_id") is not None
    suf = "_pred" if has_pred else "_gt"
    diag_id = dr.get("predicted_diagnosis_id") if has_pred else dr.get("ground_truth")
    lines = [f"Doctor's final diagnosis: {diag_id}={DISORDER_NAMES.get(diag_id, diag_id)}"]
    if not has_pred:
        lines.append("(note: doctor's raw diagnosis text didn't map to a KG disease; "
                      "checklist below is scored against ground truth as the closest reference.)")
    for grp in dr.get(f"symptom_criterion_evaluations{suf}") or []:
        matched = [_short_name(d) for d in (grp.get("matched_symptom_descriptions") or [])]
        missing = [_short_name(d) for d in (grp.get("missing_required_symptoms") or [])]
        lines.append(
            f"  [{grp.get('group')}] need {grp.get('min_count_required')}/{grp.get('relation')}: "
            f"have=[{', '.join(matched) or '-'}] missing=[{', '.join(missing) or '-'}]"
        )
    for label, key in (("duration", "duration_score"), ("functional_impairment", "functional_impairment_score"),
                        ("additional_requirements", "additional_requirements_score")):
        v = dr.get(f"{key}{suf}")
        if v is not None:
            lines.append(f"  {label}: {'met' if v >= 1.0 else 'not met'}")
    return "\n".join(lines)


RUBRIC = """Score a doctor's psychiatric intake interview on 3 dimensions, 1-5 each.
[h] Diagnostic Hypothesis Quality (every turn): does the differential (DDx) at
that point sensibly reflect info so far — no major omissions, no stale
irrelevant candidates, updates appropriately on new evidence? NOT about final
diagnosis correctness. 1=major omissions/stale candidates, no updating.
3=reasonable but some gaps or partial updating. 5=well-formed at each point,
updates cleanly on new evidence.
[q] Diagnostic Question Quality (every turn with a Q; T0 is not scored): does the question help rule out
or confirm current DDx candidates by targeting their key required symptoms,
without repeating already-resolved info? Not about how much info was
eventually gathered. 1=mostly unhelpful/repetitive. 3=somewhat useful but
unfocused or some repetition. 5=targets the sharpest current uncertainty
efficiently, no repetition.
[evidence_sufficiency] (once, whole case): does the evidence actually
gathered in the dialogue (see checklist: have= vs missing=) sufficiently
support the doctor's final diagnosis per DSM-5 (required symptoms met,
duration/functional/additional requirements met)? Mention alone isn't
enough — judge by have= vs missing= counts. 1=key evidence largely missing.
3=partially supported, notable gaps remain. 5=well supported, dialogue alone
justifies the diagnosis."""

OUTPUT_SPEC = ('Reply with ONLY this JSON, no prose: '
               '{"turns":[{"t":<int, from 1>,"h":<1-5>,"q":<1-5 or null, null if that turn has no Q>},...],'
               '"evidence_sufficiency":<1-5>}')


def build_prompt(case: CaseData) -> str:
    legend = _legend(case)
    opening, turns = policy_turns(case.transcript["turns"])
    turn_lines = [f"T0 (opening question, not scored)\nQ: {(opening or '').strip()}"]
    for t in turns:
        cs = t.get("candidate_set") or {}
        turn_lines.append(
            f"T{t['turn']}\n"
            f"A: {t['patient_response'].strip()}\n"
            f"DDx: {_tier_codes(cs)}\n"
            + (f"Q: {t['question'].strip()}" if t["question"] else "Q: (none — final turn)")
        )
    checklist = _checklist_block(case)
    return (
        f"{RUBRIC}\n\n"
        f"Disease legend: {legend}\n\n"
        f"--- Dialogue ---\n" + "\n\n".join(turn_lines) + "\n\n"
        f"--- Final diagnosis vs. DSM-5 checklist ---\n{checklist}\n\n"
        f"{OUTPUT_SPEC}"
    )


# ── LLM call ─────────────────────────────────────────────────────────────

_client = None


def _get_client():
    global _client
    if _client is None:
        from openai import OpenAI
        _client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])
    return _client


def call_llm(prompt: str) -> str:
    client = _get_client()
    try:
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.0,
        )
    except Exception as e:
        if "temperature" not in str(e).lower():
            raise
        # reasoning-style models (gpt-5/6*) only accept the default temperature.
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[{"role": "user", "content": prompt}],
        )
    return resp.choices[0].message.content or ""


_JSON_RE = re.compile(r"\{.*\}", re.DOTALL)


def parse_response(raw: str) -> dict | None:
    m = _JSON_RE.search(raw)
    if not m:
        return None
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError:
        return None


# ── Main ─────────────────────────────────────────────────────────────────

def run_case(sheet: str, profile_id: str, doctor: str) -> dict:
    case = CaseData(profile_id, doctor)
    prompt = build_prompt(case)
    raw = call_llm(prompt)
    parsed = parse_response(raw)
    result = {
        "sheet": sheet, "profile_id": profile_id, "doctor": doctor,
        "turn_numbering": TURN_NUMBERING,
        "raw_response": raw, "parsed": parsed,
    }
    out_path = OUT_DIR / f"{sheet}.json"
    out_path.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--print-prompt", metavar="SHEET", help="print the built prompt for one sheet (e.g. D001-1) and exit, no API call")
    ap.add_argument("--case", metavar="SHEET", help="run only this one sheet")
    args = ap.parse_args()

    cases = load_case_list()
    print(f"Loaded {len(cases)} cases from expert_response_2")

    if args.print_prompt:
        sheet, profile_id, doctor = next(c for c in cases if c[0] == args.print_prompt)
        case = CaseData(profile_id, doctor)
        prompt = build_prompt(case)
        print(f"\n=== Prompt for {sheet} ({profile_id} / {doctor}) — {len(prompt)} chars ===\n")
        print(prompt)
        return

    if args.case:
        cases = [c for c in cases if c[0] == args.case]

    all_results = []
    for sheet, profile_id, doctor in cases:
        print(f"  scoring {sheet} ({profile_id} / {doctor}) ...")
        try:
            result = run_case(sheet, profile_id, doctor)
        except Exception as e:
            print(f"    [error] {e}", file=sys.stderr)
            continue
        if result["parsed"] is None:
            print(f"    [warn] failed to parse JSON from LLM response", file=sys.stderr)
        all_results.append(result)

    summary_path = OUT_DIR / "llm_judge_scores.json"
    summary_path.write_text(json.dumps(all_results, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"[saved] {summary_path} ({len(all_results)} cases)")


if __name__ == "__main__":
    main()
