#!/usr/bin/env python3
"""
Korean-language version of ../generate_judge_xlsx.py.

Reuses the same 92 sampled (profile_id, doctor) cases and the same automated
metrics / diagnostic-checklist data, but renders the dialogue, patient
profile, and all section labels in Korean. Dialogue translations and the
disorder/symptom-name glossary come from pre-translated JSON produced by
translation agents (see /tmp scratchpad — not part of the repo).
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from generate_judge_xlsx import (  # noqa: E402
    CaseData, sample_cases, results_dir, analysis_dir, PROFILES_DIR,
)

TRANSLATED_DIR = Path(
    "/tmp/claude-1003/-home-jsshin-mental-simulation/59041d6a-8621-4fad-ac5c-671374a46b87/scratchpad/kor_translated"
)
OUT_DIR = Path(__file__).resolve().parent
FONT_NAME = "Malgun Gothic"

GLOSSARY = json.load(open(TRANSLATED_DIR / "_glossary_ko.json", encoding="utf-8"))
DISORDERS_KO = GLOSSARY["disorders"]
SYMPTOMS_KO = GLOSSARY["symptoms"]
GENDER_KO = GLOSSARY["gender"]
SEVERITY_KO = GLOSSARY["severity"]
EDUCATION_KO = GLOSSARY["education"]

# reverse-lookup EN "Name Words: description." -> symptom code, to translate
# diagnostic_reasoning_eval.json's matched/missing description strings
import glob as _glob  # noqa: E402

_EN_SYMPTOMS = {}
for fp in _glob.glob(str(PROFILES_DIR.parent.parent / "mentalbench/resources/knowledge_graph/EN/symptom/*.json")):
    _EN_SYMPTOMS.update(json.load(open(fp, encoding="utf-8")))
EN_DESC_TO_CODE = {
    f"{v['name'].replace('_', ' ')}: {v['description']}": k for k, v in _EN_SYMPTOMS.items()
}


def dname_ko(did: str) -> str:
    return DISORDERS_KO.get(did, did) if did else did


def symptom_desc_ko(en_desc: str) -> str:
    code = EN_DESC_TO_CODE.get(en_desc)
    if not code or code not in SYMPTOMS_KO:
        return en_desc
    s = SYMPTOMS_KO[code]
    return f"{s['name']}: {s['description']}"


class KoCaseData(CaseData):
    def __init__(self, profile_id: str, doctor: str):
        super().__init__(profile_id, doctor)
        with open(TRANSLATED_DIR / f"{profile_id}__{doctor}.json", encoding="utf-8") as f:
            self.ko = json.load(f)
        self._ko_turns = {t["turn"]: t for t in self.ko["turns"]}

    def patient_profile_text_ko(self):
        cp = self.profile["clinical_profiles"]
        pi = self.profile["patient_info"]
        feats = cp.get("sampled_features", cp)
        crit = feats.get("sampled_diagnostic_criteria", {})
        parts = [p.strip() for p in pi.get("demographics", "").split("/")]
        gender = GENDER_KO.get(parts[0], parts[0]) if len(parts) > 0 else ""
        age = parts[1] if len(parts) > 1 else ""
        employment = parts[2].replace("_", " ") if len(parts) > 2 else ""
        education = EDUCATION_KO.get(parts[3], parts[3]) if len(parts) > 3 else ""
        additional = ", ".join(crit.get("additional_requirements", []) or [])
        lines = [
            f"진단명: {dname_ko(cp.get('disease_code'))} ({cp.get('disease_code')})",
            f"인구통계: 성별 {gender} / 나이 {age} / 직업상태 {employment} / 최종학력 {education}",
            f"중증도: {SEVERITY_KO.get(pi.get('severity'), pi.get('severity'))}",
            f"스트레스 요인: {self.ko.get('stressor_event')}",
            f"증상 지속 기간: {crit.get('duration')}",
            f"추가 요건: {additional}",
            f"주호소 증상 코드: {pi.get('chief_complaint_symptom_code')}",
            "",
            f"배경 서술: {self.ko.get('stressor')}",
        ]
        return "\n".join(str(l) for l in lines)

    @staticmethod
    def _candidate_str_ko(candidate_set: dict) -> str:
        parts = []
        for tier, label in (
            ("high_likely", "고확률"),
            ("moderate_likely", "중간확률"),
            ("low_likely", "저확률"),
            ("excluded", "배제됨"),
        ):
            ids = candidate_set.get(tier) or []
            if not ids:
                continue
            names = "; ".join(f"{dname_ko(i)} ({i})" for i in ids)
            parts.append(f"{label}: {names}")
        return " | ".join(parts) if parts else "(이 시점에 기록된 감별진단 없음)"

    def turns_ko(self):
        for t in self.transcript["turns"]:
            turn = t["turn"]
            ko_t = self._ko_turns.get(turn, {})
            yield (
                turn,
                ko_t.get("doctor_question", t["doctor_question"]),
                ko_t.get("patient_response", t["patient_response"]),
                self._candidate_str_ko(t.get("candidate_set", {})),
            )

    def final_diagnosis_text_ko(self):
        eff = self.efficiency
        return (
            f"의사의 최종 진단: {dname_ko(eff.get('final_diagnosis_id'))} "
            f"({eff.get('final_diagnosis_id')}, ICD: {eff.get('final_diagnosis')})"
        )

    def checklist_rows_ko(self):
        dr = self.diagnostic_reasoning
        rows = [
            f"지속 기간 요건 — 점수: {dr.get('duration_score')}",
            f"기능적 손상 요건 — 점수: {dr.get('functional_impairment_score')}",
            f"추가 요건 — 점수: {dr.get('additional_requirements_score')}",
        ]
        for grp in dr.get("symptom_criterion_evaluations") or []:
            matched = "; ".join(symptom_desc_ko(d) for d in (grp.get("matched_symptom_descriptions") or [])) or "(없음)"
            missing = "; ".join(symptom_desc_ko(d) for d in (grp.get("missing_required_symptoms") or [])) or "(없음)"
            rows.append(
                f"체크리스트 그룹: {grp.get('group')} "
                f"(관계={grp.get('relation')}, 요구개수={grp.get('min_count_required')}, "
                f"충족개수={grp.get('valid_symptom_count')}, 커버리지={grp.get('symptom_coverage_score')})\n"
                f"  확인된 근거: {matched}\n"
                f"  누락: {missing}"
            )
        return rows


METRIC_LABELS_KO = [
    "자카드 유사도 (턴 평균)",
    "정밀도 (턴 평균)",
    "재현율 (턴 평균)",
    "정보습득점수 IAS (턴 평균)",
    "후보축소기대값 ECR (턴 평균)",
    "총 턴 수",
    "첫 확신적 후보 축소까지의 턴 수",
    "최종 진단 정확도",
    "진단적 추론 능력 (overall_score)",
]

TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True)
SECTION_FONT = Font(name=FONT_NAME, size=13, bold=True)
LABEL_FONT = Font(name=FONT_NAME, size=11, bold=True)
BODY_FONT = Font(name=FONT_NAME, size=10)
SECTION_FILL = PatternFill("solid", fgColor="E8E8E8")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def write_section_header(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT
    for col in (1, 2):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    return row + 1


def build_workbook(case: KoCaseData) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "대화문"

    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 90

    row = 1
    title_cell = ws.cell(
        row=row, column=1,
        value=f"케이스 ID: {case.profile_id}    |    의사 모델: {case.doctor}",
    )
    title_cell.font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 2

    row = write_section_header(ws, row, "자동 평가 지표")
    scores = case.scores_table()
    for label_ko, (_, value) in zip(METRIC_LABELS_KO, scores):
        ws.cell(row=row, column=1, value=label_ko).font = LABEL_FONT
        ws.cell(row=row, column=2, value=value).font = BODY_FONT
        row += 1
    row += 1

    row = write_section_header(ws, row, "환자 프로필")
    cell = ws.cell(row=row, column=1, value=case.patient_profile_text_ko())
    cell.font = BODY_FONT
    cell.alignment = WRAP_TOP
    row += 2

    row = write_section_header(ws, row, "대화문")
    transcript_start_row = row
    for turn, question, response, candidates in case.turns_ko():
        c = ws.cell(row=row, column=2, value=f"[{turn}번째 턴 · 의사 질문] {question}")
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1

        c = ws.cell(row=row, column=1, value=f"[{turn}번째 턴 · 환자 응답] {response}")
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1

        c = ws.cell(row=row, column=2, value=f"[{turn}번째 턴 · 의사의 감별진단] {candidates}")
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1
    row += 1

    row = write_section_header(ws, row, "최종 진단 및 진단 체크리스트")
    c = ws.cell(row=row, column=1, value=case.final_diagnosis_text_ko())
    c.font = LABEL_FONT
    c.alignment = WRAP_TOP
    row += 1

    for entry in case.checklist_rows_ko():
        c = ws.cell(row=row, column=1, value=entry)
        c.font = BODY_FONT
        c.alignment = WRAP_TOP
        row += 1

    ws.freeze_panes = f"A{transcript_start_row}"
    return wb


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    samples = sample_cases()
    print(f"Rendering {len(samples)} Korean workbooks")
    written = 0
    for profile_id, doctor in samples:
        try:
            case = KoCaseData(profile_id, doctor)
        except Exception as e:
            print(f"  SKIP {profile_id} / {doctor}: {e}")
            continue
        wb = build_workbook(case)
        out_path = OUT_DIR / f"judge_{profile_id}__{doctor}_kor.xlsx"
        wb.save(out_path)
        written += 1
        print(f"  wrote {out_path.name}")
    print(f"Done: {written}/{len(samples)} workbooks written to {OUT_DIR}")


if __name__ == "__main__":
    main()
