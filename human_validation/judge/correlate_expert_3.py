#!/usr/bin/env python3
"""
Expert vs. automated-metric correlation for doctor_전문가평가_{2,3}_new.xlsx.

Each workbook has one sheet per case (46 sheets), with 1-5 ratings in:
  col E (idx 4) = 진단 가설 품질 (hypothesis quality)
  col F (idx 5) = 질문 품질 (question quality)
  col G (idx 6) = 진단 근거 충분성 (evidence sufficiency)
Per-case expert score = mean of non-empty cells in each column.
Expert score is reported per expert and as the mean of the two experts.

Paired automated metrics:
  진단 가설 품질  <->  jaccard (rigid, mean over turns)
  질문 품질       <->  QTS (IAS, episode mean)
  진단 근거 충분성 <->  diagnostic_evidence_sufficiency_pred

Usage:
  MS_RUN=run_batch_20260912 python human_validation/judge/correlate_expert_3.py
"""
from __future__ import annotations

import glob
import sys as _sys
from pathlib import Path

import numpy as np
import openpyxl

_THIS_DIR = Path(__file__).resolve().parent
_REPO_ROOT = _THIS_DIR.parent.parent
for p in (_REPO_ROOT, _THIS_DIR):
    if str(p) not in _sys.path:
        _sys.path.insert(0, str(p))

from plot_expert_vs_llm_correlation import (  # noqa: E402
    authoritative_mean_ias,
    authoritative_overall_score_pred,
    _row_turn,
    _sheet_uses_policy_numbering,
    authoritative_rigid_means,
    doctor_from_kor_path,
    find_kor_file,
    icc_2_1,
    pearson,
    spearman,
    turn_ias,
    turn_jaccard_rigid,
)
from utils.turn_policy import OPENING_TURN  # noqa: E402

EXPERT_FILES = [
    sorted(glob.glob(str(_THIS_DIR / "expert_response_2" / f"doctor_*{n}_new.xlsx")))[0]
    for n in (2, 3)
]

DIMS = [
    ("hypothesis", "진단 가설 품질", 4, "jaccard"),
    ("question", "질문 품질", 5, "ias"),
    ("evidence", "진단 근거 충분성", 6, "diagnostic_evidence_sufficiency_pred"),
]


def parse_expert(path: str) -> dict[str, dict[str, float]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: dict[str, dict[str, float]] = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        header = ws.cell(row=1, column=2).value
        if not header or "/" not in str(header):
            continue
        pid = str(header).split("/")[0].strip()
        vals: dict[str, list[float]] = {key: [] for key, *_ in DIMS}
        for row in ws.iter_rows(min_row=2, values_only=True):
            for key, _label, idx, _metric in DIMS:
                if len(row) > idx and isinstance(row[idx], (int, float)):
                    vals[key].append(float(row[idx]))
        out[pid] = {k: (float(np.mean(v)) if v else None) for k, v in vals.items()}
    return out


def automated_metrics(pid: str) -> dict[str, float | None] | None:
    kor = find_kor_file(pid)
    if kor is None:
        return None
    doctor = doctor_from_kor_path(kor)
    rigid = authoritative_rigid_means(doctor, pid)
    return {
        "jaccard": rigid.get("jaccard"),
        "ias": authoritative_mean_ias(doctor, pid),
        "diagnostic_evidence_sufficiency_pred": authoritative_overall_score_pred(doctor, pid),
    }


def fmt(v: float | None) -> str:
    return f"{v:.3f}" if v is not None else "n/a"


def main() -> None:
    experts = [parse_expert(p) for p in EXPERT_FILES]
    for p, e in zip(EXPERT_FILES, experts):
        print(f"{Path(p).name}: {len(e)} cases")

    common = sorted(set(experts[0]) & set(experts[1]))
    print(f"Common cases: {len(common)}\n")

    auto = {}
    missing = []
    for pid in common:
        m = automated_metrics(pid)
        if m is None:
            missing.append(pid)
        else:
            auto[pid] = m
    if missing:
        print(f"[warn] no kor xlsx for: {missing}", file=_sys.stderr)

    print("=== Inter-expert agreement (ICC(2,1) / Pearson) ===")
    for key, label, *_ in DIMS:
        xs = [experts[0][pid][key] for pid in common]
        ys = [experts[1][pid][key] for pid in common]
        pairs = [(x, y) for x, y in zip(xs, ys) if x is not None and y is not None]
        xs2, ys2 = [a for a, _ in pairs], [b for _, b in pairs]
        icc, _ = icc_2_1(xs2, ys2)
        r, n = pearson(xs2, ys2)
        print(f"  {label:<12} ICC={fmt(icc)}  r={fmt(r)}  (n={n})")

    print("\n=== Expert vs. automated metric ===")
    header = f"{'expert score':<18}{'metric':<40}{'rater':<14}{'n':>4}{'Pearson r':>12}{'Spearman rho':>14}"
    print(header)
    print("-" * len(header))
    for key, label, _idx, metric in DIMS:
        raters = {
            "Expert 2": lambda pid: experts[0][pid][key],
            "Expert 3": lambda pid: experts[1][pid][key],
            "Mean(2,3)": lambda pid: (
                float(np.mean([v for v in (experts[0][pid][key], experts[1][pid][key]) if v is not None]))
                if any(v is not None for v in (experts[0][pid][key], experts[1][pid][key])) else None
            ),
        }
        for rater, getter in raters.items():
            xs, ys = [], []
            for pid in auto:
                e = getter(pid)
                a = auto[pid].get(metric)
                if e is None or a is None:
                    continue
                xs.append(e)
                ys.append(a)
            if len(xs) < 3:
                continue
            r, n = pearson(xs, ys)
            rho, _ = spearman(xs, ys)
            print(f"{label:<18}{metric:<40}{rater:<14}{n:>4}{fmt(r):>12}{fmt(rho):>14}")
        print()

    print(f"[saved] {draw_scatter(auto, experts, common)}")
    turn_level(common)


def expert_turn_ratings(path: str) -> dict[str, dict[str, dict[int, float]]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: dict[str, dict[str, dict[int, float]]] = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        header = ws.cell(row=1, column=2).value
        if not header or "/" not in str(header):
            continue
        pid = str(header).split("/")[0].strip()
        policy = _sheet_uses_policy_numbering(ws)
        hyp: dict[int, float] = {}
        qst: dict[int, float] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            turn = _row_turn(row[2] if len(row) > 2 else None, policy)
            if turn is None or turn == OPENING_TURN:
                continue
            if len(row) > 4 and isinstance(row[4], (int, float)):
                hyp[turn] = float(row[4])
            if len(row) > 5 and isinstance(row[5], (int, float)):
                qst[turn] = float(row[5])
        out[pid] = {"hyp": hyp, "qst": qst}
    return out


def turn_level(common: list[str]) -> None:
    per_expert = [expert_turn_ratings(p) for p in EXPERT_FILES]
    pairs = {"hyp": ("jaccard", turn_jaccard_rigid), "qst": ("ias", turn_ias)}
    labels = {"hyp": "진단 가설 품질 vs Jaccard (rigid)", "qst": "질문 품질 vs QTS (IAS)"}
    print("\n=== Turn-level correlation (pooled turns, expert mean of 2) ===")
    for key, (metric, getter) in pairs.items():
        xs, ys = [], []
        for pid in common:
            doctor_kor = find_kor_file(pid)
            if doctor_kor is None:
                continue
            doctor = doctor_from_kor_path(doctor_kor)
            auto_by_turn = getter(doctor, pid)
            for t, a in auto_by_turn.items():
                vals = [e[pid][key].get(t) for e in per_expert if pid in e]
                vals = [v for v in vals if v is not None]
                if not vals:
                    continue
                xs.append(a)
                ys.append(float(np.mean(vals)))
        r, n = pearson(xs, ys)
        rho, _ = spearman(xs, ys)
        print(f"{labels[key]:<36} n_turns={n:<5} Pearson r={fmt(r)}  Spearman rho={fmt(rho)}")


def draw_scatter(auto: dict[str, dict], experts: list[dict], common: list[str]) -> Path:
    import matplotlib.font_manager as fm
    import matplotlib.pyplot as plt

    cjk = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
    fm.fontManager.addfont(cjk)
    plt.rcParams["font.family"] = [fm.FontProperties(fname=cjk).get_name(), "DejaVu Sans"]

    fig, axes = plt.subplots(1, 3, figsize=(16, 5.4))
    metric_labels = {
        "jaccard": "Jaccard (rigid)",
        "ias": "QTS (IAS)",
        "diagnostic_evidence_sufficiency_pred": "DES (_pred)",
    }
    for ax, (key, label, _idx, metric) in zip(axes, DIMS):
        xs, ys = [], []
        for pid in common:
            a = auto.get(pid, {}).get(metric)
            vals = [e[pid][key] for e in experts if e[pid][key] is not None]
            if a is None or not vals:
                continue
            xs.append(a)
            ys.append(float(np.mean(vals)))
        xs, ys = np.array(xs), np.array(ys)
        r, _n = pearson(xs.tolist(), ys.tolist())
        rho, _ = spearman(xs.tolist(), ys.tolist())

        ax.scatter(xs, ys, s=46, color="#0072B2", alpha=0.75, edgecolors="white", linewidths=0.6, zorder=3)
        slope, intercept = np.polyfit(xs, ys, 1)
        grid = np.linspace(xs.min(), xs.max(), 50)
        ax.plot(grid, slope * grid + intercept, color="#D55E00", linewidth=1.6, zorder=2)

        ax.set_xlabel(metric_labels[metric], fontsize=13)
        ax.set_ylabel(f"{label} (expert mean, 1-5)", fontsize=13)
        ax.set_ylim(0.8, 5.2)
        ax.tick_params(labelsize=11)
        ax.grid(color="#E4EBF5", linewidth=0.8, zorder=0)
        ax.spines[["top", "right"]].set_visible(False)
        ax.set_title(f"r = {fmt(r)},  ρ = {fmt(rho)}  (n={len(xs)})", fontsize=13)

    fig.tight_layout()
    out = _THIS_DIR / "expert_vs_metrics_scatter.png"
    fig.savefig(out, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out


if __name__ == "__main__":
    main()
